#!/usr/bin/env python3
"""
Export Pre-Approvals to Excel with Images
Exports pre-approval data from Oracle database to Excel with embedded images
"""

import sys
import getpass
import argparse
from datetime import datetime
import io

try:
    import oracledb
except ImportError:
    print("Error: oracledb module not found.")
    print("Please install it using: pip install oracledb")
    sys.exit(1)

try:
    import openpyxl
    from openpyxl.drawing.image import Image as XLImage
    from openpyxl.utils import get_column_letter
except ImportError:
    print("Error: openpyxl module not found.")
    print("Please install it using: pip install openpyxl")
    sys.exit(1)

try:
    from PIL import Image
except ImportError:
    print("Error: Pillow module not found.")
    print("Please install it using: pip install Pillow")
    sys.exit(1)


def get_pre_approvals_data(connection):
    """
    Fetch pre-approval data from the database

    Args:
        connection: Oracle database connection

    Returns:
        list: List of tuples containing pre-approval data
    """
    query = """
    select a.pre_approval_id, a.status, a.pre_approval_request_date,
           a.trouble_cd_text, a.correction_text, b.comment_text
    from dbo.pre_approvals a, dbo.pre_approval_comments b
    where a.pre_approval_id = b.pre_approval_id
    and a.PRE_APPROVAL_REQUEST_DATE > (Sysdate-90)
    and a.status in ('ACCEPTED', 'REJECTED')
    and a.LABOR_OPERATION_CD IN
     ('7020230','7021960','7021970','7022350','7022890','7022930','7022940',
     '7023350','7023358','7023370','7023390','7023410','7023430','7023510',
     '7023520','7023550','7023600','7023750','7023770','7023870','7022780',
     '7022810','7022690','7023130','7024150','7024210','7024430','7024940',
     '7025090','7025250','7025260', '7025280','7025350','7025360','7025390',
    '7025510','7025600','7025620','7024720','7028090 ','7028110 ')
    order by a.PRE_APPROVAL_REQUEST_DATE desc, b.sequence_id
    """

    cursor = connection.cursor()
    cursor.execute(query)
    results = cursor.fetchall()
    cursor.close()

    return results


def get_images_for_pre_approval(connection, pre_approval_id):
    """
    Fetch images for a specific pre-approval ID

    Args:
        connection: Oracle database connection
        pre_approval_id: The pre-approval ID to fetch images for

    Returns:
        list: List of tuples containing (image_data, filename)
    """
    query = """
    SELECT c.PRE_APPROVAL_ATTACHMENT, b.FILENAME
    from PRE_APPROVAL_META b, pre_approval_attachment_data3 c
    WHERE b.PRE_APPROVAL_ATTACHMENT_ID = c.PRE_APPROVAL_ATTACHMENT_ID
    AND UPPER(b.FILENAME) LIKE '%JPG%'
    and b.PRE_APPROVAL_ATTACHMENT_ID = :pre_approval_id
    """

    cursor = connection.cursor()
    cursor.execute(query, {'pre_approval_id': pre_approval_id})
    results = cursor.fetchall()
    cursor.close()

    return results


def create_excel_with_images(data, connection, output_file='pre_approvals_with_images.xlsx'):
    """
    Create Excel file with data and embedded images

    Args:
        data: List of tuples containing pre-approval data
        connection: Oracle database connection
        output_file: Output Excel file path
    """
    # Create a new workbook and select active sheet
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Pre-Approvals"

    # Define headers based on the query columns
    headers = [
        'Pre-Approval ID',
        'Status',
        'Pre-Approval Request Date',
        'Trouble Code Text',
        'Correction Text',
        'Comment Text',
        'Images'
    ]

    # Write headers
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num)
        cell.value = header
        cell.font = openpyxl.styles.Font(bold=True)

    # Set column widths
    ws.column_dimensions['A'].width = 15  # Pre-Approval ID
    ws.column_dimensions['B'].width = 12  # Status
    ws.column_dimensions['C'].width = 20  # Pre-Approval Request Date
    ws.column_dimensions['D'].width = 30  # Trouble Code Text
    ws.column_dimensions['E'].width = 30  # Correction Text
    ws.column_dimensions['F'].width = 40  # Comment Text
    ws.column_dimensions['G'].width = 60  # Images

    # Default row height
    default_row_height = 15

    # Write data and embed images
    current_row = 2
    print(f"\nProcessing {len(data)} records...")

    for record in data:
        pre_approval_id, status, request_date, trouble_text, correction_text, comment_text = record

        # Write data to cells
        ws.cell(row=current_row, column=1, value=pre_approval_id)
        ws.cell(row=current_row, column=2, value=status)
        ws.cell(row=current_row, column=3, value=request_date)
        ws.cell(row=current_row, column=4, value=trouble_text)
        ws.cell(row=current_row, column=5, value=correction_text)
        ws.cell(row=current_row, column=6, value=comment_text)

        # Fetch and embed images for this pre-approval
        try:
            images = get_images_for_pre_approval(connection, pre_approval_id)

            if images:
                print(f"  Found {len(images)} image(s) for Pre-Approval ID {pre_approval_id}")

                # Calculate row height based on number of images
                # Each image will be about 100 pixels tall
                row_height = max(default_row_height, len(images) * 75)
                ws.row_dimensions[current_row].height = row_height

                # Position for placing images
                image_col = 7  # Column G
                image_offset = 0

                for img_data, filename in images:
                    try:
                        # Convert BLOB to image
                        if img_data:
                            # Read BLOB data
                            blob_data = img_data.read() if hasattr(img_data, 'read') else img_data

                            # Create PIL Image from blob
                            pil_image = Image.open(io.BytesIO(blob_data))

                            # Resize image to fit in cell (max width 400px, max height 100px)
                            max_width = 400
                            max_height = 100
                            pil_image.thumbnail((max_width, max_height), Image.Resampling.LANCZOS)

                            # Save to BytesIO object
                            img_buffer = io.BytesIO()
                            pil_image.save(img_buffer, format='PNG')
                            img_buffer.seek(0)

                            # Create openpyxl image
                            xl_image = XLImage(img_buffer)

                            # Position the image in the cell
                            cell_address = f"{get_column_letter(image_col)}{current_row}"
                            xl_image.anchor = cell_address

                            # Adjust image position vertically for multiple images
                            if image_offset > 0:
                                # Offset in pixels
                                xl_image.anchor = f"{get_column_letter(image_col)}{current_row}"
                                # Note: openpyxl doesn't support fine-grained vertical offset within same cell
                                # Images will stack, but may overlap. Consider using separate rows if needed.

                            ws.add_image(xl_image)

                            # Add filename as hyperlink or text
                            cell = ws.cell(row=current_row, column=image_col)
                            if cell.value:
                                cell.value = f"{cell.value}\n{filename}"
                            else:
                                cell.value = filename

                            image_offset += 105  # Offset for next image

                    except Exception as e:
                        print(f"    Warning: Could not process image {filename}: {e}")

            else:
                # No images found
                ws.row_dimensions[current_row].height = default_row_height

        except Exception as e:
            print(f"  Error fetching images for Pre-Approval ID {pre_approval_id}: {e}")
            ws.row_dimensions[current_row].height = default_row_height

        current_row += 1

    # Save the workbook
    wb.save(output_file)
    print(f"\nExcel file created successfully: {output_file}")
    print(f"Total records: {len(data)}")


def main(username=None, password=None, output_file='pre_approvals_with_images.xlsx'):
    """
    Main function to export pre-approvals to Excel with images

    Args:
        username (str, optional): Database username
        password (str, optional): Database password
        output_file (str): Output Excel file path
    """
    # Database connection parameters
    hostname = "dcmpgwm04-cl.edc.nam.gm.com"
    port = 1521
    service_name = "gwmtkpd_srv.edc.nam.gm.com"

    print("Pre-Approvals Export to Excel")
    print(f"Hostname: {hostname}")
    print(f"Service Name: {service_name}")
    print("-" * 50)

    # Get credentials
    if username is None:
        username = input("Enter username: ").strip()

    if not username:
        print("Error: Username cannot be empty")
        return

    if password is None:
        password = getpass.getpass("Enter password: ")

    if not password:
        print("Error: Password cannot be empty")
        return

    # Construct DSN
    dsn = oracledb.makedsn(hostname, port, service_name=service_name)

    print("\nConnecting to Oracle database...")

    try:
        # Connect to database
        connection = oracledb.connect(
            user=username,
            password=password,
            dsn=dsn
        )

        print("Connected successfully!")

        # Fetch pre-approvals data
        print("\nFetching pre-approvals data...")
        data = get_pre_approvals_data(connection)

        if not data:
            print("No data found matching the criteria.")
            connection.close()
            return

        print(f"Found {len(data)} records")

        # Create Excel with images
        print("\nCreating Excel file with images...")
        create_excel_with_images(data, connection, output_file)

        # Close connection
        connection.close()
        print("\nDatabase connection closed.")
        print("Export completed successfully!")

    except oracledb.Error as error:
        print(f"\nDatabase Error: {error}")
        error_obj, = error.args
        if hasattr(error_obj, 'code'):
            print(f"Error Code: {error_obj.code}")
        if hasattr(error_obj, 'message'):
            print(f"Error Message: {error_obj.message}")

    except Exception as error:
        print(f"\nUnexpected error: {error}")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(
        description="Export pre-approvals to Excel with embedded images",
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog="""
Examples:
  # Interactive mode (prompts for credentials):
  python export_pre_approvals.py

  # With username only (prompts for password):
  python export_pre_approvals.py -u myusername

  # With username, password, and custom output file:
  python export_pre_approvals.py -u myusername -p mypassword -o output.xlsx
        """
    )
    parser.add_argument(
        '-u', '--username',
        type=str,
        help='Database username (will prompt if not provided)'
    )
    parser.add_argument(
        '-p', '--password',
        type=str,
        help='Database password (will prompt if not provided)'
    )
    parser.add_argument(
        '-o', '--output',
        type=str,
        default='pre_approvals_with_images.xlsx',
        help='Output Excel file path (default: pre_approvals_with_images.xlsx)'
    )

    args = parser.parse_args()

    main(username=args.username, password=args.password, output_file=args.output)
